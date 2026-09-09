#!/usr/bin/env python3
"""
render_status_update.py — Amcor-branded storytelling status update (.docx)

Usage:
    python3 render_status_update.py spec.json output.docx

The spec is a JSON object. See SKILL.md for the full schema and BLOCK REFERENCE.
Everything is optional except `header` and `sections`.
"""
import json
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- Amcor palette
DARK = '00395A'    # Amcor Dark Blue — dominant
LBLUE = '00A1DE'   # Amcor Light Blue — accents, rules, numbers
GREEN = '00A551'   # Amcor Green — good news, settled items
ORANGE = 'E98300'  # Amcor Orange — exposure / risk figures only
GREY = '5A6872'
BODY = '333F48'
LIGHTBG = 'F2F7FA'
RULE = 'D6E2EA'

NAMED = {'dark': DARK, 'lblue': LBLUE, 'green': GREEN, 'orange': ORANGE,
         'grey': GREY, 'body': BODY}


def col(v, default=BODY):
    if not v:
        return default
    return NAMED.get(str(v).lower(), str(v).lstrip('#'))


# ---------------------------------------------------------------- doc scaffold
doc = Document()
_s = doc.sections[0]
_s.page_width, _s.page_height = Inches(8.5), Inches(11)
for _a in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(_s, _a, Inches(0.75))

_st = doc.styles['Normal']
_st.font.name = 'Arial'
_st.font.size = Pt(10)
_st.font.color.rgb = RGBColor.from_string(BODY)
_st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
_st.paragraph_format.space_after = Pt(6)
_st.paragraph_format.line_spacing = 1.15


# ---------------------------------------------------------------- primitives
def sub(el, tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    el.append(e)
    return e


def run(par, text, size=10, bold=False, italic=False, color=BODY, caps=False, spacing=None):
    r = par.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.all_caps = caps
    if spacing:
        rPr = r._element.get_or_add_rPr()
        e = OxmlElement('w:spacing')
        e.set(qn('w:val'), str(spacing))
        rPr.append(e)
    return r


def para(container, before=0, after=6, align=None, indent=None, keep=False, line=1.15):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    pf.keep_with_next = keep
    if align:
        pf.alignment = align
    if indent:
        pf.left_indent = Inches(indent)
    return p


def bottom_border(p, color=LBLUE, size=12, space=6):
    sub(sub(p._p.get_or_add_pPr(), 'w:pBdr'), 'w:bottom',
        val='single', sz=size, space=space, color=color)


def spacer(pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def shade(cell, fill):
    sub(cell._tc.get_or_add_tcPr(), 'w:shd', val='clear', color='auto', fill=fill)


def cell_borders(cell, top=None, bottom=None, left=None, right=None):
    b = sub(cell._tc.get_or_add_tcPr(), 'w:tcBorders')
    for side, spec in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        if spec is None:
            sub(b, 'w:' + side, val='nil')
        else:
            sub(b, 'w:' + side, val='single', sz=spec[0], space='0', color=spec[1])


def cell_margins(cell, t=90, b=90, l=120, r=120):
    m = sub(cell._tc.get_or_add_tcPr(), 'w:tcMar')
    for side, v in (('top', t), ('bottom', b), ('left', l), ('right', r)):
        sub(m, 'w:' + side, w=v, type='dxa')


def clear_cell(cell):
    cell.text = ''
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.line_spacing = 1.15
    return p0


def new_table(cols_in):
    t = doc.add_table(rows=1, cols=len(cols_in))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    sub(t._tbl.tblPr, 'w:tblLayout', type='fixed')
    grid = t._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), cols_in):
        gc.set(qn('w:w'), str(int(w * 1440)))
    for r in t.rows:
        sub(r._tr.get_or_add_trPr(), 'w:cantSplit')
        for c, w in zip(r.cells, cols_in):
            c.width = Inches(w)
    return t


def add_row(t, cols_in):
    r = t.add_row()
    sub(r._tr.get_or_add_trPr(), 'w:cantSplit')
    for c, w in zip(r.cells, cols_in):
        c.width = Inches(w)
    return r


def rich(par, value, base_size=10, base_color=BODY, base_italic=False):
    """value: a plain string, or a list of segments.
    Segment = "text" | {"t": "text", "b": true, "i": true, "color": "dark"}"""
    if isinstance(value, str):
        run(par, value, size=base_size, color=base_color, italic=base_italic)
        return
    for seg in value:
        if isinstance(seg, str):
            run(par, seg, size=base_size, color=base_color, italic=base_italic)
        else:
            run(par, seg.get('t', ''),
                size=seg.get('size', base_size),
                bold=seg.get('b', False),
                italic=seg.get('i', base_italic),
                color=col(seg.get('color'), base_color))


# ---------------------------------------------------------------- blocks
def b_h1(blk):
    p = para(doc, before=9, after=5, keep=True)
    if blk.get('num'):
        run(p, str(blk['num']) + '  ', size=12, bold=True, color=LBLUE)
    run(p, blk['text'], size=13, bold=True, color=DARK)
    bottom_border(p, color=RULE, size=6, space=5)


def b_h2(blk):
    p = para(doc, before=11, after=5, keep=True)
    run(p, blk['text'], size=10.5, bold=True, color=DARK)


def b_lead(blk):
    """One grey line of framing under a heading. Keep it to a single sentence."""
    p = para(doc, after=7)
    rich(p, blk['text'], base_color=GREY)


def b_say(blk):
    """THE core block: what Damian actually says out loud, in quotes."""
    lines = blk['lines'] if isinstance(blk.get('lines'), list) else [blk.get('text', '')]
    t = new_table([7.0])
    c = t.rows[0].cells[0]
    shade(c, LIGHTBG)
    cell_borders(c, left=(24, col(blk.get('accent'), LBLUE)))
    cell_margins(c, t=95, b=95, l=170, r=170)
    p = clear_cell(c)
    for i, ln in enumerate(lines):
        tgt = p if i == 0 else c.add_paragraph()
        tgt.paragraph_format.space_after = Pt(0 if i == len(lines) - 1 else 7)
        tgt.paragraph_format.space_before = Pt(0)
        tgt.paragraph_format.line_spacing = 1.15
        run(tgt, '“', size=10, italic=True, color=DARK)
        rich(tgt, ln, base_color=DARK, base_italic=True)
        run(tgt, '”', size=10, italic=True, color=DARK)
    spacer(3)


def b_bullets(blk):
    for it in blk['items']:
        p = para(doc, after=blk.get('after', 4), indent=0.24)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        run(p, '•  ', size=10, bold=True, color=LBLUE)
        if isinstance(it, dict) and 'label' in it:
            run(p, it['label'] + ' ', size=10, bold=True, color=DARK)
            rich(p, it.get('text', ''))
        else:
            rich(p, it)


def b_kpis(blk):
    items = blk['items']
    w = 7.0 / len(items)
    t = new_table([w] * len(items))
    for c, it in zip(t.rows[0].cells, items):
        shade(c, LIGHTBG)
        cell_borders(c, right=(8, 'FFFFFF'))
        cell_margins(c, t=100, b=100, l=140, r=140)
        p = clear_cell(c)
        p.paragraph_format.space_after = Pt(1)
        run(p, it['value'], size=14, bold=True, color=col(it.get('color'), DARK))
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.space_before = Pt(0)
        run(p2, it['label'], size=7.5, color=GREY)
    spacer(6)


def b_evidence(blk):
    """3-column evidence table: impact | evidence | value at stake.
    Never state a value without the evidence that produced it."""
    cols = blk.get('widths', [2.15, 3.05, 1.80])
    hdr = blk.get('header', ['Impacto al negocio', 'Evidencia', 'Valor en juego'])
    t = new_table(cols)
    sub(t.rows[0]._tr.get_or_add_trPr(), 'w:tblHeader')
    for i, (c, label) in enumerate(zip(t.rows[0].cells, hdr)):
        cell_borders(c, bottom=(6, RULE))
        cell_margins(c, t=40, b=70, l=0 if i == 0 else 120, r=0 if i == len(cols) - 1 else 120)
        p = clear_cell(c)
        if i == len(cols) - 1:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(p, label, size=7.5, bold=True, color=GREY, caps=True, spacing=20)
    for r in blk['rows']:
        cs = add_row(t, cols).cells
        for i, c in enumerate(cs):
            cell_borders(c, bottom=(6, RULE))
            cell_margins(c, t=62, b=62, l=0 if i == 0 else 120, r=0 if i == len(cols) - 1 else 120)
        p = clear_cell(cs[0])
        run(p, r['impact'], size=9.5, bold=True, color=DARK)
        p = clear_cell(cs[1])
        rich(p, r.get('evidence', ''), base_size=8.5)
        p = clear_cell(cs[2])
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(1)
        vcolor = col(r.get('color'), DARK)
        big = r.get('value', '')
        run(p, big, size=11.5 if r.get('color') else 9.5, bold=True, color=vcolor)
        if r.get('note'):
            p2 = cs[2].add_paragraph()
            p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.space_before = Pt(0)
            run(p2, r['note'], size=7, color=GREY)
    spacer(6)


def b_status_cols(blk):
    """Settled vs still-open. This is the honesty block — use it whenever
    something could be mistaken for a decision that has not been made."""
    left, right = blk['left'], blk['right']
    t = new_table([3.5, 3.5])
    for c, side, default in ((t.rows[0].cells[0], left, GREEN),
                             (t.rows[0].cells[1], right, ORANGE)):
        accent = col(side.get('color'), default)
        shade(c, LIGHTBG)
        cell_borders(c, left=(20, accent), right=(10, 'FFFFFF'))
        cell_margins(c, t=100, b=100, l=150, r=150)
        p = clear_cell(c)
        p.paragraph_format.space_after = Pt(5)
        run(p, side['title'], size=7.5, bold=True, color=accent, caps=True, spacing=20)
        for it in side['items']:
            q = c.add_paragraph()
            q.paragraph_format.space_before = Pt(0)
            q.paragraph_format.space_after = Pt(4)
            q.paragraph_format.line_spacing = 1.12
            q.paragraph_format.left_indent = Inches(0.15)
            q.paragraph_format.first_line_indent = Inches(-0.15)
            run(q, '•  ', size=9, bold=True, color=accent)
            if isinstance(it, dict) and 'label' in it:
                run(q, it['label'] + ' ', size=9, bold=True, color=DARK)
                rich(q, it.get('text', ''), base_size=9)
            else:
                rich(q, it, base_size=9)
        c.paragraphs[-1].paragraph_format.space_after = Pt(0)
    spacer(6)


def b_qa(blk):
    """Anticipated questions and the answer, in Damian's voice."""
    for it in blk['items']:
        p = para(doc, before=7, after=2, keep=True)
        run(p, it['q'], size=10, bold=True, color=DARK)
        p = para(doc, after=3, indent=0.16)
        run(p, '“', size=9.5, italic=True, color=BODY)
        rich(p, it['a'], base_size=9.5, base_italic=True)
        run(p, '”', size=9.5, italic=True, color=BODY)


def b_note(blk):
    p = para(doc, before=blk.get('before', 3), after=blk.get('after', 6))
    rich(p, blk['text'])


def b_pagebreak(_blk):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def b_subheader(blk):
    """Divider title for the annex pages."""
    if blk.get('eyebrow'):
        p = para(doc, after=3)
        run(p, blk['eyebrow'], size=8, bold=True, color=LBLUE, spacing=30)
    p = para(doc, after=2)
    run(p, blk['title'], size=17, bold=True, color=DARK)
    bottom_border(p, color=LBLUE, size=12, space=7)
    spacer(10)


def b_spacer(blk):
    spacer(blk.get('pts', 6))


BLOCKS = {
    'h1': b_h1, 'h2': b_h2, 'lead': b_lead, 'say': b_say, 'bullets': b_bullets,
    'kpis': b_kpis, 'evidence': b_evidence, 'status_cols': b_status_cols,
    'qa': b_qa, 'note': b_note, 'pagebreak': b_pagebreak,
    'subheader': b_subheader, 'spacer': b_spacer,
}


# ---------------------------------------------------------------- main
def main():
    spec = json.load(open(sys.argv[1], encoding='utf-8'))
    out = sys.argv[2]

    h = spec.get('header', {})
    p = para(doc, after=3)
    parts = h.get('eyebrow', ['BUSINESS UPDATE'])
    if isinstance(parts, str):
        parts = [parts]
    for i, seg in enumerate(parts):
        if i:
            run(p, '   ·   ', size=8, color=RULE)
        run(p, seg, size=8, bold=True, color=LBLUE, spacing=30)

    if h.get('title'):
        p = para(doc, after=3)
        run(p, h['title'], size=20, bold=True, color=DARK)
    if h.get('subtitle'):
        p = para(doc, after=2)
        run(p, h['subtitle'], size=9, color=GREY)
        bottom_border(p, color=LBLUE, size=12, space=7)
    if h.get('meta'):
        p = para(doc, before=4, after=8)
        run(p, h['meta'], size=8, italic=True, color=GREY)

    for blk in spec.get('sections', []):
        fn = BLOCKS.get(blk.get('type'))
        if fn is None:
            raise SystemExit('Unknown block type: %r. Valid: %s'
                             % (blk.get('type'), ', '.join(sorted(BLOCKS))))
        fn(blk)

    doc.save(out)
    print('OK ->', out)


if __name__ == '__main__':
    main()
